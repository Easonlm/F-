"""Create the editable Word manuscript from the copy-ready Markdown paper."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "问题一论文正文.md"
OUTPUT = HERE / "问题一论文正文.docx"


def set_font(run, name="SimSun", size=10.5, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if run._element.get_or_add_rPr().rFonts is not None:
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def math_text(value):
    s = value
    replace = {
        r"\epsilon": "ε", r"\alpha": "α", r"\Delta": "Δ", r"\phi": "φ", r"\mathbf": "",
        r"\mathrm": "", r"\operatorname": "", r"\widehat": "", r"\bar": "",
        r"\sum": "∑", r"\max": "max", r"\min": "min", r"\in": "∈", r"\ne": "≠",
        r"\geq": "≥", r"\leq": "≤", r"\ldots": "…", r"\qquad": "  ", r"\quad": " ",
        r"\left": "", r"\right": "", r"\|": "|", r"\{": "{", r"\}": "}",
        r"\,": " ", r"\,": " ",
    }
    for old, new in replace.items():
        s = s.replace(old, new)
    # Keep formulas editable as plain text; the Markdown retains the LaTeX source.
    for _ in range(4):
        s = re.sub(r"_\{([^{}]+)\}", r"_\1", s)
    s = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", s)
    s = s.replace("\\", "")
    s = s.replace("{", "").replace("}", "")
    return s


def display_equation(value):
    if "z_{jk}=" in value:
        return "z_jk = clip[(t_jk - a_k)/(b_k - a_k), 0, 1]"
    if "S_{j,\\mathrm{edu}}=" in value:
        return "S_j,edu = [z_j,fineweb + z_j,reasoning + z_j,professionalism + z_j,qurater + (z_j,dsir_books + z_j,dsir_wiki + z_j,dsir_math)/3]/5"
    if "Q_j=" in value:
        return "Q_j = (S_j,edu + S_j,read + S_j,clean + S_j,structure)/4,  Q_j ∈ [0,1]"
    if "x_{si}=" in value:
        return "x_si = ln[(p_si + ε)/(p_sr + ε)],  i ≠ r"
    if "widehat" in value and "min" in value:
        return "L̂_s = b + B·std[φ(x_s)];  min Σ_s ||L_s − L̂_s||² + α||B||²_F"
    return math_text(value).strip("$")


def add_inline(paragraph, value, size=10.5, first_bold=False):
    value = re.sub(r"\$([^$]+)\$", lambda m: math_text(m.group(1)), value)
    value = value.replace("`", "")
    pieces = value.split("**")
    for i, piece in enumerate(pieces):
        if not piece:
            continue
        r = paragraph.add_run(piece)
        set_font(r, size=size, bold=(i % 2 == 1) or first_bold)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def border_table(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D9D9D9")
        borders.append(e)
    tbl_pr.append(borders)


def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(value))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tc_pr.append(mar)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def set_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit")
    tr_pr.append(flag)


def add_table(doc, block):
    rows = []
    for line in block:
        cells = [x.strip() for x in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", x) for x in cells):
            continue
        rows.append(cells)
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    border_table(table)
    if ncols == 6:
        widths = [3.0, 2.5, 2.1, 2.1, 2.1, 2.4]
    elif ncols == 4:
        widths = [3.3, 2.2, 3.0, 6.0]
    else:
        widths = [15.0 / ncols] * ncols
    for r_idx, row in enumerate(rows):
        tr = table.rows[r_idx]
        set_no_split(tr)
        if r_idx == 0:
            repeat_header(tr)
        for c_idx, value in enumerate(row):
            cell = tr.cells[c_idx]
            cell.width = Cm(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value, size=8.5, first_bold=(r_idx == 0))
            if r_idx == 0:
                shade(cell, "2F4058")
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif r_idx % 2 == 0:
                shade(cell, "F2F5F8")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.7)
    for name, size, before, after in [("Heading 1", 13, 14, 7), ("Heading 2", 11, 10, 5)]:
        style = doc.styles[name]
        style.font.name = "SimHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)
    title = doc.styles["Title"]
    title.font.name = "SimHei"
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
    title.paragraph_format.space_after = Pt(16)

    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, line[2:], size=16, first_bold=True)
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("$$") and line.endswith("$$") and len(line) > 4:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(display_equation(line[2:-2]))
            set_font(r, name="Cambria Math", size=9.5)
        elif line == "$$":
            equation = []
            i += 1
            while i < len(lines) and lines[i].strip() != "$$":
                equation.append(lines[i].strip())
                i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(display_equation(" ".join(equation)))
            set_font(r, name="Cambria Math", size=9.5)
        elif line.startswith("| "):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            add_table(doc, block)
            continue
        elif line.startswith("!["):
            match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
            if match:
                path = HERE / match.group(1)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(path), width=Cm(14.8))
        elif line.startswith("*图") and line.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, line.strip("*"), size=9)
        elif line.startswith("**表"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, line, size=9, first_bold=True)
        else:
            p = doc.add_paragraph()
            is_reference = line.startswith("[1]") or line.startswith("[2]") or line.startswith("[3]")
            if is_reference:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.left_indent = Cm(0.7)
                p.paragraph_format.first_line_indent = Cm(-0.7)
                p.paragraph_format.space_after = Pt(3)
            add_inline(p, line, size=9 if is_reference else 10.5)
        i += 1
    doc.core_properties.title = "问题一 数据质量评价 质量冲突消解与领域配比建模"
    doc.core_properties.subject = "数学建模竞赛问题一论文正文"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
